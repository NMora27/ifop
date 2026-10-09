"""Genera BPPALS con los parámetros entregados por main.py."""

from pathlib import Path
import warnings

from plantillas_bppals import cargar_plantilla_bppals, guardar_bppals_con_diseno

from lectura import (
    buscar_desembarque,
    leer_frecuencias_colocap,
    leer_matriz_talla_edad,
    leer_parametros_regresion_iqr,
)
from metodos_matematicos import calcular_a1_a2
from openpyxl import load_workbook


# ============================================================
# PROCESO PRINCIPAL
def generar_bppals(especie, sexo, anio, zona, area, flota_desembarque, archivo_matriz,
                   archivo_regresion, archivo_colocap, archivo_desembarque,
                   salida_machos, salida_hembras, talla_min, talla_max, paso_talla,
                   edad_min, edad_max, hoja_colocap):
    ESPECIE = especie
    SEXO = sexo
    ANIO = anio
    ZONA = zona
    AREA = area
    FLOTA_DESEMBARQUE = flota_desembarque
    ARCHIVO_MATRIZ = Path(archivo_matriz)
    ARCHIVO_REGRESION = Path(archivo_regresion)
    ARCHIVO_COLOCAP = Path(archivo_colocap)
    ARCHIVO_DESEMBARQUE = Path(archivo_desembarque)
    TALLA_MIN = talla_min
    TALLA_MAX = talla_max
    PASO_TALLA = paso_talla
    EDAD_MIN = edad_min
    EDAD_MAX = edad_max
    FILA_CLAVE_INICIO = 75
    FILA_CLAVE_PLANTILLA_FIN = 126
    FILA_TOTAL_CLAVE = 127
    FILA_PROB_INICIO = 9
    FILA_PROB_PLANTILLA_FIN = 60
    HOJA_COLOCAP = hoja_colocap
    FILA_COLOCAP_INICIO = 11
    FILA_COLOCAP_FIN = 62
    COL_EDAD_INICIO = 7

    if SEXO == "Hembras":
        SALIDA = Path(salida_hembras)
    elif SEXO == "Machos":
        SALIDA = Path(salida_machos)
    else:
        raise ValueError('SEXO debe ser "Machos" o "Hembras"')

    SALIDA.parent.mkdir(parents=True, exist_ok=True)

    for archivo in (ARCHIVO_MATRIZ, ARCHIVO_REGRESION, ARCHIVO_COLOCAP,
                    ARCHIVO_DESEMBARQUE):
        if not archivo.exists():
            raise FileNotFoundError(f"No existe el archivo: {archivo}")

    # ============================================================
    # 1. PREPROCESAMIENTO
    
    # ------------------------------------------------------------
    # 1.1 LEER REGRESIONES IQR
    
    a_machos,b_machos,a_hembras,b_hembras = leer_parametros_regresion_iqr(
        ARCHIVO_REGRESION)
    
    
    # ------------------------------------------------------------
    # 1.2 LEER MATRIZ TALLA-EDAD
    
    # ------------------------------------------------------------
    # 1.3 PREPARAR MATRIZ DE CONTEOS
    
    tallas = list(range(TALLA_MIN, TALLA_MAX + 1, PASO_TALLA))
    edades = list(range(EDAD_MIN, EDAD_MAX + 1))
    datos = leer_matriz_talla_edad(
        ARCHIVO_MATRIZ, SEXO, TALLA_MIN, TALLA_MAX, PASO_TALLA, EDAD_MIN, EDAD_MAX)
    
    
    # ------------------------------------------------------------
    # 1.4 LEER FRECUENCIAS PARA COLOCAP
    
    total_ejemplares,n_machos,n_hembras,frecuencias_colocap = leer_frecuencias_colocap(
        ARCHIVO_COLOCAP, HOJA_COLOCAP, PASO_TALLA)
    
    
    # ============================================================
    # 3. CÁLCULOS AUXILIARES
    
    # ------------------------------------------------------------
    # 3.1 PARÁMETROS a1 Y a2
    
    a1_machos,a2_machos = calcular_a1_a2(b_machos)
    a1_hembras,a2_hembras = calcular_a1_a2(b_hembras)
    
    
    # ------------------------------------------------------------
    # 3.2 DESEMBARQUE
    
    desembarque = buscar_desembarque(ARCHIVO_DESEMBARQUE,ESPECIE,FLOTA_DESEMBARQUE)
    
    
    # ============================================================
    # 4. PREPARAR ARCHIVO FINAL
    
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Cannot parse header or footer.*")
        warnings.filterwarnings("ignore", message="wmf image format is not supported.*")
        wb = load_workbook(cargar_plantilla_bppals(SEXO))
    
    HOJAS_CONSERVAR = ("Clave","Colocap","Comp.Capt.","Matrices",
                       "Graf.","Tabla","Distr.Capt.Nº","Tabla Print")
    
    for nombre in wb.sheetnames.copy():
        if nombre not in HOJAS_CONSERVAR:
            del wb[nombre]
    
    for nombre in HOJAS_CONSERVAR:
        if nombre not in wb.sheetnames:
            raise ValueError(f"La plantilla no contiene la hoja {nombre}")
    
    ws_clave = wb["Clave"]
    ws_colocap = wb["Colocap"]
    ws_comp = wb["Comp.Capt."]
    ws_matrices = wb["Matrices"]
    ws_graf = wb["Graf."]
    ws_tabla = wb["Tabla"]
    ws_distr = wb["Distr.Capt.Nº"]
    ws_print = wb["Tabla Print"]
    
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    
    
    # ============================================================
    # 5. GENERAR HOJA CLAVE
    
    # ------------------------------------------------------------
    # 5.1 METADATOS
    ws_clave["B71"] = f"Especie: {ESPECIE.capitalize()}"
    ws_clave["K71"] = f"Sexo: {SEXO}"
    ws_clave["T71"] = f"Zona: {ZONA}"
    ws_clave["AB71"] = f"Fecha: {ANIO}"
    
    # ------------------------------------------------------------
    # 5.2 CONFIGURAR MATRIZ
    COL_EDAD_FIN = COL_EDAD_INICIO + len(edades) - 1
    FILA_CLAVE_FIN = FILA_CLAVE_INICIO + len(tallas) - 1
    FILA_PROB_FIN = FILA_PROB_INICIO + len(tallas) - 1
    
    if FILA_CLAVE_FIN > FILA_CLAVE_PLANTILLA_FIN:
        raise ValueError("La cantidad de tallas supera el espacio disponible en Clave")
    if FILA_PROB_FIN > FILA_PROB_PLANTILLA_FIN:
        raise ValueError("La cantidad de tallas supera el espacio de probabilidades")
    
    # Limpiar valores antiguos de la plantilla
    for fila in range(FILA_CLAVE_INICIO,FILA_CLAVE_PLANTILLA_FIN + 1):
        for columna in range(1,COL_EDAD_FIN + 1):
            ws_clave.cell(fila,columna).value = None
    
    for fila in range(FILA_PROB_INICIO,FILA_PROB_PLANTILLA_FIN + 1):
        for columna in range(2,COL_EDAD_FIN + 1):
            ws_clave.cell(fila,columna).value = None
    
    # ------------------------------------------------------------
    # 5.3 CLAVE EDAD-TALLA
    
    for edad in range(EDAD_MIN,EDAD_MAX):
        columna = COL_EDAD_INICIO + (edad - EDAD_MIN)
        ws_clave.cell(74,columna).value = edad
    
    ws_clave.cell(74,COL_EDAD_FIN).value = f"{EDAD_MAX}+"
    
    for i,talla in enumerate(tallas):
        fila = FILA_CLAVE_INICIO + i
    
        ws_clave.cell(fila,2).value = talla
        ws_clave.cell(fila,3).value = "-"
        ws_clave.cell(fila,4).value = talla + PASO_TALLA - 1
    
        for edad in range(EDAD_MIN,EDAD_MAX):
            columna = COL_EDAD_INICIO + (edad - EDAD_MIN)
            valor = datos[talla][edad]
            ws_clave.cell(fila,columna).value = None if valor == 0 else valor
    
        valor_40 = datos[talla][EDAD_MAX]
        ws_clave.cell(fila,COL_EDAD_FIN).value = None if valor_40 == 0 else valor_40
    
        col_inicio = ws_clave.cell(fila,COL_EDAD_INICIO).column_letter
        col_fin = ws_clave.cell(fila,COL_EDAD_FIN).column_letter
        ws_clave.cell(fila,6).value = f"=SUM({col_inicio}{fila}:{col_fin}{fila})"
    
    # Actualizar fila Total
    for columna in [6] + list(range(COL_EDAD_INICIO,COL_EDAD_FIN + 1)):
        letra = ws_clave.cell(FILA_TOTAL_CLAVE,columna).column_letter
        ws_clave.cell(FILA_TOTAL_CLAVE,columna).value = \
            f"=SUM({letra}{FILA_CLAVE_INICIO}:{letra}{FILA_CLAVE_FIN})"
    
    # ------------------------------------------------------------
    # 5.4 PROBABILIDAD DE PERTENENCIA
    
    for i in range(len(tallas)):
        fila_prob = FILA_PROB_INICIO + i
        fila_clave = FILA_CLAVE_INICIO + i
    
        ws_clave.cell(fila_prob,2).value = f"=B{fila_clave}"
        ws_clave.cell(fila_prob,3).value = "-"
        ws_clave.cell(fila_prob,4).value = f"=D{fila_clave}"
    
        for columna in range(COL_EDAD_INICIO,COL_EDAD_FIN + 1):
            letra = ws_clave.cell(fila_prob,columna).column_letter
            ws_clave.cell(fila_prob,columna).value = \
                f'=IF($F{fila_clave}=0,0,{letra}{fila_clave}/$F{fila_clave})'
    
        col_inicio = ws_clave.cell(fila_prob,COL_EDAD_INICIO).column_letter
        col_fin = ws_clave.cell(fila_prob,COL_EDAD_FIN).column_letter
        ws_clave.cell(fila_prob,6).value = \
            f"=SUM({col_inicio}{fila_prob}:{col_fin}{fila_prob})"
    
    # ------------------------------------------------------------
    # 5.5 ENCABEZADOS TABLA SUPERIOR
    for edad in range(EDAD_MIN,EDAD_MAX):
        columna = COL_EDAD_INICIO + (edad - EDAD_MIN)
        ws_clave.cell(8,columna).value = f"={ws_clave.cell(74,columna).coordinate}"
    
    ws_clave.cell(8,COL_EDAD_FIN).value = \
        f"={ws_clave.cell(74,COL_EDAD_FIN).coordinate}"
    
    
    # ============================================================
    # 6. GENERAR HOJA COLOCAP
    
    # ------------------------------------------------------------
    # 6.1 DESEMBARQUE
    ws_colocap["C8"] = desembarque
    
    # ------------------------------------------------------------
    # 6.2 PARÁMETROS LONGITUD-PESO
    ws_colocap["C6"] = a_machos
    ws_colocap["C7"] = b_machos
    ws_colocap["D6"] = a_hembras
    ws_colocap["D7"] = b_hembras
    
    
    # ------------------------------------------------------------
    # 6.3 PARÁMETROS a1 Y a2
    ws_colocap["C68"] = a1_machos
    ws_colocap["C69"] = a2_machos
    ws_colocap["D68"] = a1_hembras
    ws_colocap["D69"] = a2_hembras
    
    
    # ------------------------------------------------------------
    # 6.4 FRECUENCIAS POR CLASE DE TALLA
    
    for fila in range(FILA_COLOCAP_INICIO,FILA_COLOCAP_FIN + 1):
        marca_clase = ws_colocap.cell(fila,2).value
    
        try:
            marca_clase = int(marca_clase)
        except (TypeError,ValueError):
            continue
    
        if marca_clase in frecuencias_colocap:
            ws_colocap.cell(fila,3).value = frecuencias_colocap[marca_clase]["Machos"]
            ws_colocap.cell(fila,4).value = frecuencias_colocap[marca_clase]["Hembras"]
        else:
            ws_colocap.cell(fila,3).value = None
            ws_colocap.cell(fila,4).value = None
    
    
    # ============================================================
    # 7. GENERAR HOJA COMP.CAPT.
    
    FILA_COMP_INICIO = 8
    FILA_COMP_PLANTILLA_FIN = 59
    FILA_COMP_FIN = FILA_COMP_INICIO + len(tallas) - 1
    COL_COMP_EDAD_INICIO = 6
    COL_COMP_EDAD_FIN = COL_COMP_EDAD_INICIO + len(edades) - 1
    
    # ------------------------------------------------------------
    # 7.1 METADATOS
    ws_comp["A4"] = "=Clave!$B$71"
    ws_comp["J4"] = "=Clave!$K$71"
    ws_comp["S4"] = "=Clave!$T$71"
    ws_comp["AA4"] = "=Clave!$AB$71"
    
    # ------------------------------------------------------------
    # 7.2 LIMPIAR TABLA ANTIGUA
    for fila in range(FILA_COMP_INICIO,FILA_COMP_PLANTILLA_FIN + 1):
        for columna in range(1,COL_COMP_EDAD_FIN + 1):
            ws_comp.cell(fila,columna).value = None
    
    # ------------------------------------------------------------
    # 7.3 ENCABEZADOS DE EDAD
    for i in range(len(edades)):
        col_comp = COL_COMP_EDAD_INICIO + i
        col_clave = COL_EDAD_INICIO + i
        letra_clave = ws_clave.cell(74,col_clave).column_letter
        ws_comp.cell(7,col_comp).value = f"=Clave!{letra_clave}74"
    
    # ------------------------------------------------------------
    # 7.4 RELACIÓN MARCA DE CLASE - COLOCAP
    filas_marcas = {}
    
    for fila in range(FILA_COLOCAP_INICIO,FILA_COLOCAP_FIN + 1):
        marca = ws_colocap.cell(fila,2).value
    
        try:
            marca = int(marca)
        except (TypeError,ValueError):
            continue
    
        filas_marcas[marca] = fila
    
    # ------------------------------------------------------------
    # 7.5 COMPOSICIÓN DE CAPTURA
    
    col_frec = "D" if SEXO == "Hembras" else "C"
    celda_numero = "$F$6" if SEXO == "Hembras" else "$E$6"
    
    for i,talla in enumerate(tallas):
        fila_comp = FILA_COMP_INICIO + i
        fila_clave = FILA_CLAVE_INICIO + i
        fila_prob = FILA_PROB_INICIO + i
        marca_clase = talla + PASO_TALLA // 2
    
        if marca_clase not in filas_marcas:
            raise ValueError(
                f"No se encontró la marca de clase {marca_clase} en Colocap")
    
        fila_colocap = filas_marcas[marca_clase]
    
        ws_comp.cell(fila_comp,1).value = f"=Clave!B{fila_clave}"
        ws_comp.cell(fila_comp,2).value = "-"
        ws_comp.cell(fila_comp,3).value = f"=Clave!D{fila_clave}"
    
        ws_comp.cell(fila_comp,5).value = \
            f"=(Colocap!{col_frec}{fila_colocap}/Colocap!{col_frec}$64)*Colocap!{celda_numero}"
    
        for j in range(len(edades)):
            col_comp = COL_COMP_EDAD_INICIO + j
            col_clave = COL_EDAD_INICIO + j
            letra_clave = ws_clave.cell(fila_prob,col_clave).column_letter
    
            ws_comp.cell(fila_comp,col_comp).value = \
                f"=(Clave!{letra_clave}{fila_prob}*$E{fila_comp})"
    
    # ------------------------------------------------------------
    # 7.6 TOTALES
    
    ws_comp["E61"] = f"=SUM(E{FILA_COMP_INICIO}:E{FILA_COMP_FIN})"
    
    for columna in range(COL_COMP_EDAD_INICIO,COL_COMP_EDAD_FIN + 1):
        letra = ws_comp.cell(61,columna).column_letter
        ws_comp.cell(61,columna).value = \
            f"=SUM({letra}{FILA_COMP_INICIO}:{letra}{FILA_COMP_FIN})"
    
    ws_comp["E62"] = \
        f"=SUM({ws_comp.cell(61,COL_COMP_EDAD_INICIO).column_letter}61:" \
        f"{ws_comp.cell(61,COL_COMP_EDAD_FIN).column_letter}61)"
    
    # ------------------------------------------------------------
    # 7.7 PORCENTAJES
    for columna in range(COL_COMP_EDAD_INICIO,COL_COMP_EDAD_FIN + 1):
        letra = ws_comp.cell(63,columna).column_letter
        ws_comp.cell(63,columna).value = f"=({letra}61*100)/$E$61"
    
    
    # ============================================================
    # 8. GENERAR HOJA MATRICES
    
    FILA_MAT_INICIO = 9
    FILA_MAT_FIN = FILA_MAT_INICIO + len(tallas) - 1
    FILA_VAR_INICIO = 70
    FILA_VAR_FIN = FILA_VAR_INICIO + len(tallas) - 1
    COL_MAT_EDAD_INICIO = 3
    COL_MAT_EDAD_FIN = COL_MAT_EDAD_INICIO + len(edades) - 1
    
    # ------------------------------------------------------------
    # 8.1 METADATOS
    for fila in (5,66):
        ws_matrices.cell(fila,1).value = "=Clave!$B$71"
        ws_matrices.cell(fila,7).value = "=Clave!$K$71"
        ws_matrices.cell(fila,16).value = "=Clave!$T$71"
        ws_matrices.cell(fila,24).value = "=Clave!$AB$71"
    
    # ------------------------------------------------------------
    # 8.2 ENCABEZADOS DE EDAD
    for i in range(len(edades)):
        col_mat = COL_MAT_EDAD_INICIO + i
        col_clave = COL_EDAD_INICIO + i
        letra_clave = ws_clave.cell(74,col_clave).column_letter
    
        ws_matrices.cell(8,col_mat).value = f"=Clave!{letra_clave}74"
        ws_matrices.cell(69,col_mat).value = f"={ws_matrices.cell(8,col_mat).coordinate}"
    
    # ------------------------------------------------------------
    # 8.3 LIMPIAR MATRICES ANTIGUAS
    for fila in range(9,61):
        for columna in range(1,COL_MAT_EDAD_FIN + 1):
            ws_matrices.cell(fila,columna).value = None
    
    for fila in range(70,122):
        for columna in range(1,COL_MAT_EDAD_FIN + 1):
            ws_matrices.cell(fila,columna).value = None
    
    # ------------------------------------------------------------
    # 8.4 MATRIZ TALLA × FRECUENCIA
    
    for i,talla in enumerate(tallas):
        fila_mat = FILA_MAT_INICIO + i
        fila_comp = FILA_COMP_INICIO + i
        marca_clase = talla + PASO_TALLA // 2
        col_fin = ws_matrices.cell(fila_mat,COL_MAT_EDAD_FIN).column_letter
    
        ws_matrices.cell(fila_mat,1).value = marca_clase
        ws_matrices.cell(fila_mat,2).value = f"=SUM(C{fila_mat}:{col_fin}{fila_mat})"
    
        for j in range(len(edades)):
            col_mat = COL_MAT_EDAD_INICIO + j
            col_comp = COL_COMP_EDAD_INICIO + j
            letra_comp = ws_comp.cell(fila_comp,col_comp).column_letter
    
            ws_matrices.cell(fila_mat,col_mat).value = \
                f"=$A{fila_mat}*'Comp.Capt.'!{letra_comp}{fila_comp}"
    
    # Totales
    ws_matrices["A61"] = "Total"
    
    for columna in range(2,COL_MAT_EDAD_FIN + 1):
        letra = ws_matrices.cell(61,columna).column_letter
        ws_matrices.cell(61,columna).value = \
            f"=SUM({letra}{FILA_MAT_INICIO}:{letra}{FILA_MAT_FIN})"
    
    # ------------------------------------------------------------
    # 8.5 MATRIZ PARA CÁLCULO DE VARIANZA
    
    for i in range(len(tallas)):
        fila_var = FILA_VAR_INICIO + i
        fila_mat = FILA_MAT_INICIO + i
        fila_comp = FILA_COMP_INICIO + i
        col_fin = ws_matrices.cell(fila_var,COL_MAT_EDAD_FIN).column_letter
    
        ws_matrices.cell(fila_var,1).value = f"=A{fila_mat}"
        ws_matrices.cell(fila_var,2).value = f"=SUM(C{fila_var}:{col_fin}{fila_var})"
    
        for j in range(len(edades)):
            col_mat = COL_MAT_EDAD_INICIO + j
            col_comp = COL_COMP_EDAD_INICIO + j
            letra_comp = ws_comp.cell(fila_comp,col_comp).column_letter
    
            ws_matrices.cell(fila_var,col_mat).value = \
                f"=$A{fila_var}^2*'Comp.Capt.'!{letra_comp}{fila_comp}"
    
    # Totales
    ws_matrices["A122"] = "Total"
    
    for columna in range(2,COL_MAT_EDAD_FIN + 1):
        letra = ws_matrices.cell(122,columna).column_letter
        ws_matrices.cell(122,columna).value = \
            f"=SUM({letra}{FILA_VAR_INICIO}:{letra}{FILA_VAR_FIN})"
    
    
    # ============================================================
    # 9. GENERAR HOJA GRAF.
    
    # ------------------------------------------------------------
    # 9.1 TÍTULOS Y METADATOS
    ws_graf["A1"] = "Composición del desembarque en número (%)"
    ws_graf["G1"] = f"Composición del desembarque en número (%) de {ESPECIE.lower()}"
    ws_graf["G2"] = "=Clave!$T$71"
    ws_graf["J2"] = "=Clave!$K$71"
    ws_graf["L2"] = "=Clave!$AB$71"
    
    # ------------------------------------------------------------
    # 9.2 DATOS PARA EL GRÁFICO
    ws_graf["B3"] = "GE"
    ws_graf["C3"] = "%"
    
    for i,edad in enumerate(edades):
        fila = 4 + i
        col_comp = COL_COMP_EDAD_INICIO + i
        letra_comp = ws_comp.cell(63,col_comp).column_letter
    
        ws_graf.cell(fila,2).value = f"{EDAD_MAX}+" if edad == EDAD_MAX else edad
        ws_graf.cell(fila,3).value = f"='Comp.Capt.'!{letra_comp}$63"
    
    # Cálculos auxiliares existentes en la plantilla
    ws_graf["D10"] = "=SUM(C9:C10)"
    ws_graf["D20"] = "=SUM(C19:C20)"
    ws_graf["D25"] = "=SUM(C23:C25)"
    ws_graf["E25"] = "=D10+D20+D25"
    
    
    # ============================================================
    # 10. GENERAR HOJA TABLA
    
    FILA_TABLA_INICIO = 8
#    FILA_TABLA_FIN = FILA_TABLA_INICIO + len(tallas) - 1
    FILA_TABLA_PLANTILLA_FIN = 59
    
    # ------------------------------------------------------------
    # 10.1 TÍTULO
    desembarque_txt = f"{desembarque:,.3f}"
    desembarque_txt = desembarque_txt.replace(",","X").replace(".",",").replace("X",".")
    
    ws_tabla["A3"] = (
        f"Composición del desembarque (D) en número de individuos por grupo de edad "
        f"de {ESPECIE.lower()}, {SEXO.lower()}. {ZONA}, {AREA}, {ANIO} "
        f"(Desembarque total= {desembarque_txt} t).")
    
    # ------------------------------------------------------------
    # 10.2 ENCABEZADOS DE EDAD
    for i in range(len(edades)):
        col_tabla = COL_COMP_EDAD_INICIO + i
        col_clave = COL_EDAD_INICIO + i
        letra_clave = ws_clave.cell(74,col_clave).column_letter
        ws_tabla.cell(7,col_tabla).value = f"=Clave!{letra_clave}74"
    
    # ------------------------------------------------------------
    # 10.3 LIMPIAR TABLA ANTIGUA
    for fila in range(FILA_TABLA_INICIO,FILA_TABLA_PLANTILLA_FIN + 1):
        for columna in range(1,COL_COMP_EDAD_FIN + 1):
            ws_tabla.cell(fila,columna).value = None
    
    # ------------------------------------------------------------
    # 10.4 COMPOSICIÓN TALLA-EDAD
    
    for i in range(len(tallas)):
        fila_tabla = FILA_TABLA_INICIO + i
        fila_comp = FILA_COMP_INICIO + i
        ws_tabla.cell(fila_tabla,1).value = f"='Comp.Capt.'!A{fila_comp}"
        ws_tabla.cell(fila_tabla,2).value = "-"
        ws_tabla.cell(fila_tabla,3).value = f"='Comp.Capt.'!C{fila_comp}"
    
        for columna in range(5,COL_COMP_EDAD_FIN + 1):
            letra = ws_tabla.cell(fila_tabla,columna).column_letter
            ws_tabla.cell(fila_tabla,columna).value = \
                f'=IF(\'Comp.Capt.\'!{letra}{fila_comp}=0," ",\'Comp.Capt.\'!{letra}{fila_comp})'
    
    # ------------------------------------------------------------
    # 10.5 RESULTADOS RESUMIDOS
    filas_resumen = {61: "Total", 63: "Porcentaje", 65: "Talla prom. (cm)",
                     67: "Varianza", 69: "Peso prom. (g)"}
    
    # Limpiar filas separadoras
    for fila in (60,62,64,66,68):
        for columna in range(1,COL_COMP_EDAD_FIN + 1):
            ws_tabla.cell(fila,columna).value = None
    
    for fila,texto in filas_resumen.items():
        ws_tabla.cell(fila,1).value = texto
    
        for columna in range(5,COL_COMP_EDAD_FIN + 1):
            letra = ws_tabla.cell(fila,columna).column_letter
            ws_tabla.cell(fila,columna).value = \
                f'=IF(\'Comp.Capt.\'!{letra}{fila}=0," ",\'Comp.Capt.\'!{letra}{fila})'
    
    
    # ============================================================
    # 11. GENERAR HOJA DISTR.CAPT.Nº
    
    FILA_DISTR_INICIO = 4
    FILA_DISTR_FIN = FILA_DISTR_INICIO + len(tallas) - 1
    
    # ------------------------------------------------------------
    # 11.1 ENCABEZADOS
    ws_distr["A3"] = "Intervalo de Talla"
    ws_distr["E1"] = "Nº individuos"
    ws_distr["F1"] = None
    ws_distr["G1"] = "Nº individuos (miles)"
    ws_distr["H1"] = None
    ws_distr["E2"] = None
    ws_distr["F2"] = None
    ws_distr["E3"] = "Machos"
    ws_distr["F3"] = "Hembras"
    ws_distr["G3"] = "Machos"
    ws_distr["H3"] = "Hembras"
    
    # ------------------------------------------------------------
    # 11.2 LIMPIAR DATOS ANTIGUOS
    for fila in range(4,46):
        for columna in range(1,9):
            ws_distr.cell(fila,columna).value = None
    
    # ------------------------------------------------------------
    # 11.3 DISTRIBUCIÓN POR TALLA
    
    for i,talla in enumerate(tallas):
        fila = FILA_DISTR_INICIO + i
        fila_tabla = FILA_TABLA_INICIO + i
        marca_clase = talla + PASO_TALLA // 2
    
        if marca_clase not in filas_marcas:
            raise ValueError(
                f"No se encontró la marca de clase {marca_clase} en Colocap")
    
        fila_colocap = filas_marcas[marca_clase]
    
        # Intervalo de talla
        ws_distr.cell(fila,1).value = f"=Tabla!A{fila_tabla}"
        ws_distr.cell(fila,2).value = "-"
        ws_distr.cell(fila,3).value = f"=Tabla!C{fila_tabla}"
        # Número de individuos por sexo
        ws_distr.cell(fila,5).value = \
            f"=(Colocap!C{fila_colocap}/Colocap!C$64)*Colocap!$E$6"
        ws_distr.cell(fila,6).value = \
            f"=(Colocap!D{fila_colocap}/Colocap!D$64)*Colocap!$F$6"
    
        # Número expresado en miles
        ws_distr.cell(fila,7).value = f"=E{fila}/1000"
        ws_distr.cell(fila,8).value = f"=F{fila}/1000"
    
    # ------------------------------------------------------------
    # 11.4 TOTALES
    ws_distr["E47"] = f"=SUM(E{FILA_DISTR_INICIO}:E{FILA_DISTR_FIN})"
    ws_distr["F47"] = f"=SUM(F{FILA_DISTR_INICIO}:F{FILA_DISTR_FIN})"
    ws_distr["G47"] = f"=SUM(G{FILA_DISTR_INICIO}:G{FILA_DISTR_FIN})"
    ws_distr["H47"] = f"=SUM(H{FILA_DISTR_INICIO}:H{FILA_DISTR_FIN})"
    
    
    # ============================================================
    # 12. GENERAR HOJA TABLA PRINT
    
    FILA_PRINT_INICIO = 8
#    FILA_PRINT_FIN = FILA_PRINT_INICIO + len(tallas) - 1
    
    # ------------------------------------------------------------
    # 12.1 TÍTULO
    ws_print["A3"] = "=Tabla!A3"
    
    # ------------------------------------------------------------
    # 12.2 ENCABEZADOS
    ws_print["A6"] = "Tallas (cm)"
    ws_print["E6"] = "Frec."
    ws_print["F6"] = "G r u p o s   d e   e d a d"
    
    for i in range(len(edades)):
        columna = COL_COMP_EDAD_INICIO + i
        letra = ws_tabla.cell(7,columna).column_letter
        ws_print.cell(7,columna).value = f"=Tabla!{letra}7"
    
    # ------------------------------------------------------------
    # 12.3 LIMPIAR TABLA ANTIGUA
    for fila in range(8,51):
        for columna in range(1,COL_COMP_EDAD_FIN + 1):
            ws_print.cell(fila,columna).value = None
    
    # ------------------------------------------------------------
    # 12.4 COMPOSICIÓN TALLA-EDAD
    
    for i in range(len(tallas)):
        fila_print = FILA_PRINT_INICIO + i
        fila_tabla = FILA_TABLA_INICIO + i
        ws_print.cell(fila_print,1).value = f"=Tabla!A{fila_tabla}"
        ws_print.cell(fila_print,2).value = f"=Tabla!B{fila_tabla}"
        ws_print.cell(fila_print,3).value = f"=Tabla!C{fila_tabla}"
    
        for columna in range(5,COL_COMP_EDAD_FIN + 1):
            letra = ws_print.cell(fila_print,columna).column_letter
            ws_print.cell(fila_print,columna).value = \
                f"=Tabla!{letra}{fila_tabla}"
    
    # ------------------------------------------------------------
    # 12.5 LIMPIAR CLASE SOBRANTE
    for columna in range(1,COL_COMP_EDAD_FIN + 1):
        ws_print.cell(50,columna).value = None
    
    # ------------------------------------------------------------
    # 12.6 RESÚMENES
    filas_print = { 52: 61,     # Total
                    54: 63,     # Porcentaje
                    57: 65,     # Talla promedio
                    59: 67,     # Varianza
                    61: 69 }    # Peso promedio
    
    for fila_print,fila_tabla in filas_print.items():
        ws_print.cell(fila_print,1).value = f"=Tabla!A{fila_tabla}"
    
        for columna in range(5,COL_COMP_EDAD_FIN + 1):
            letra = ws_print.cell(fila_print,columna).column_letter
            ws_print.cell(fila_print,columna).value = \
                f"=Tabla!{letra}{fila_tabla}"
    
    
    # ============================================================
    # 13. GUARDAR ARCHIVO
    
    guardar_bppals_con_diseno(wb, SEXO, SALIDA)
    print(f"Archivo generado: {SALIDA}")
    

    return SALIDA

