import re
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


# ============================================================
# DETECCIÓN DE ENCODING
def detectar_encoding_dbf(ruta_dbf):
    ldid_map = {0x01: "cp437", 0x02: "cp850", 0x03: "cp1252", 0x04: "mac_roman",
                0x08: "cp865", 0x09: "cp437", 0x0A: "cp850", 0x0B: "cp437",
                0x0D: "cp437", 0x0E: "cp850", 0x0F: "cp437", 0x10: "cp850",
                0x11: "cp437", 0x12: "cp850", 0x13: "cp932", 0x14: "cp850",
                0x15: "cp437", 0x16: "cp850", 0x17: "cp865", 0x18: "cp437",
                0x19: "cp437", 0x1A: "cp850", 0x1B: "cp437", 0x1C: "cp863",
                0x1D: "cp437", 0x1E: "cp852", 0x1F: "cp852", 0x22: "cp852",
                0x23: "cp852", 0x24: "cp860", 0x25: "cp850", 0x26: "cp866",
                0x37: "cp850", 0x40: "cp852", 0x4D: "cp936", 0x4E: "cp949",
                0x4F: "cp950", 0x50: "cp874", 0x57: "cp1252", 0x58: "cp1252",
                0x59: "cp1252", 0x64: "cp852", 0x65: "cp866", 0x66: "cp865",
                0x67: "cp861", 0x6A: "cp737", 0x6B: "cp857", 0x78: "cp950",
                0x79: "cp949", 0x7A: "cp936", 0x7B: "cp932", 0x7C: "cp874",
                0x7D: "cp1255", 0x7E: "cp1256", 0x86: "cp737", 0x87: "cp852",
                0x88: "cp857", 0xC8: "cp1250", 0xC9: "cp1251", 0xCA: "cp1254",
                0xCB: "cp1253", 0xCC: "cp1257"}
    ruta = Path(ruta_dbf)

    try:
        with ruta.open("rb") as archivo:
            archivo.seek(29)
            byte_ldid = archivo.read(1)
    except OSError:
        return "cp1252"

    if byte_ldid:
        encoding = ldid_map.get(byte_ldid[0])
        if encoding:
            return encoding

    try:
        from dbfread import DBF

        tabla = DBF(str(ruta), load=False)
        if tabla.encoding:
            return tabla.encoding
    except (OSError, UnicodeError, ValueError):
        pass

    return "cp1252"

def detectar_encoding_csv(ruta_csv, bytes_muestra=100_000):
    ruta = Path(ruta_csv)

    with ruta.open("rb") as archivo:
        datos = archivo.read(bytes_muestra)

    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            datos.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue

    return "latin-1"

# ============================================================
# LECTURAS GENERALES
def leer_dbf(ruta, encoding=None):
    from dbfread import DBF

    encoding = encoding or detectar_encoding_dbf(ruta)
    tabla = DBF(str(ruta), encoding=encoding, ignore_missing_memofile=True)
    datos = pd.DataFrame(iter(tabla))
    datos.columns = [col.strip().upper().replace(" ", "_") for col in datos.columns]

    for col in datos.select_dtypes(include=["object"]).columns:
        datos[col] = datos[col].astype(str).str.strip()

    for col in datos.columns:
        if col not in ("PERIODO", "ESCALA", "AGRUPACION"):
            datos[col] = pd.to_numeric(datos[col], errors="coerce")

    return datos

def extraer_tokens(valor):
    texto = unicodedata.normalize("NFD", str(valor).upper())
    texto = "".join(caracter for caracter in texto
                    if unicodedata.category(caracter) != "Mn")
    return re.findall(r"[A-Z]+|\d+", texto)

def validar_archivo_palangre(archivo, anio):
    archivo = Path(archivo)
    anio = str(anio)
    if not archivo.exists():
        raise FileNotFoundError(f"No existe el archivo Palangre: {archivo}")
    if archivo.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError(
            f"El archivo Palangre debe ser Excel (.xlsx o .xlsm): {archivo.name}")
    if not {"PALANGRE", anio}.issubset(extraer_tokens(archivo.stem)):
        raise ValueError(
            f"El nombre '{archivo.name}' debe incluir Palangre y el año {anio}")

    return archivo

def leer_excel_palangre(archivo, columnas, especie, anio):
    archivo = validar_archivo_palangre(archivo, anio)
    codigo_especie = str(especie)
    anio_corto = str(anio)[-2:]

    with pd.ExcelFile(archivo) as libro:
        hojas = [hoja for hoja in libro.sheet_names
                 if codigo_especie in hoja.upper()
                 and "PAL" in hoja.upper()
                 and anio_corto in hoja]

        if not hojas:
            disponibles = ", ".join(libro.sheet_names)
            raise ValueError(
                f"No se encontró una hoja con {codigo_especie}, Pal y {anio_corto} "
                f"en {archivo.name}. Hojas disponibles: {disponibles}")
        if len(hojas) > 1:
            raise ValueError(
                f"Se encontraron varias hojas Palangre para especie {codigo_especie}: "
                f"{', '.join(hojas)}")

        datos = pd.read_excel(libro, sheet_name=hojas[0])

    n_original = len(datos)
    datos = datos[columnas].copy()
    datos["LONGITUD_ESPECIMEN"] = datos["LONGITUD_ESPECIMEN"].astype(float)
    datos["PESO_ESPECIMEN"] = datos["PESO_ESPECIMEN"].astype(float)
    return datos, n_original

def leer_datos_talla_edad(archivo, columnas=None, *, anio=None, especie=None, zona=None):
    """Lee el formato histórico o normaliza el estándar utilizado desde 2026."""
    datos = pd.read_excel(archivo)
    if anio is not None and int(anio) >= 2026:
        if especie is None:
            raise ValueError("Debe indicar ESPECIE_CODIGO para la matriz desde 2026")
        especie = int(especie)
        if especie in (37, 27, 54):
            columna_edad = "EDAD"
        elif especie in (1, 2, 3, 4, 6):
            columna_edad = "EDAD_CRONOLOGICA"
        else:
            raise ValueError(f"No está definida la columna de edad para la especie {especie}")
        requeridas = ["ANIO_MUESTRO", "COS_ESPECIE", "ZONA", "LONGITUD_ESPECIMEN",
                      columna_edad, "SEXO"]
        faltantes = [col for col in requeridas if col not in datos.columns]
        if faltantes:
            raise ValueError(f"Faltan columnas del formato talla-edad desde 2026: {', '.join(faltantes)}")
        datos = datos[requeridas].copy()
        for col in requeridas:
            if col != "ZONA":
                datos[col] = pd.to_numeric(datos[col], errors="coerce")
        datos = datos[(datos["ANIO_MUESTRO"] == int(anio)) &
                      (datos["COS_ESPECIE"] == especie)].copy()
        if zona is not None and str(zona).strip():
            # Acepta códigos numéricos de Excel (1.0) y zonas expresadas como texto.
            zona_texto = str(zona).strip()
            zona_numero = pd.to_numeric(zona_texto, errors="coerce")
            if pd.notna(zona_numero):
                filtro = pd.to_numeric(datos["ZONA"], errors="coerce") == zona_numero
            else:
                filtro = datos["ZONA"].astype(str).str.strip() == zona_texto
            datos = datos[filtro].copy()
        datos = datos.rename(columns={"LONGITUD_ESPECIMEN": "LONGITUD_DEL_PEZ_(CM)",
                                      columna_edad: "Nº_DE_ANILLOS"})
        return datos[["LONGITUD_DEL_PEZ_(CM)", "Nº_DE_ANILLOS", "SEXO"]].dropna()

    columnas = columnas or ["LONGITUD_DEL_PEZ_(CM)", "Nº_DE_ANILLOS",
                            "CÓDIGO_DE_PESQUERÍA", "SEXO"]
    faltantes = [col for col in columnas if col not in datos.columns]
    if faltantes:
        raise ValueError(f"Faltan columnas del formato talla-edad histórico: {', '.join(faltantes)}")
    datos = datos[columnas].copy()

    for col in columnas:
        datos[col] = pd.to_numeric(datos[col], errors="coerce")

    return datos.dropna()

# ============================================================
# LECTURAS PARA BPPALSH
def leer_parametros_regresion_iqr(archivo, encoding=None):
    encoding = encoding or detectar_encoding_csv(archivo)
    datos = pd.read_csv(archivo, sep=";", decimal=",", encoding=encoding)
    datos["_METODO"] = datos["METODO"].astype(str).str.strip().str.upper()
    datos["_GRUPO"] = datos["GRUPO"].astype(str).str.strip().str.upper()
    datos_iqr = datos[datos["_METODO"] == "IQR"].copy()
    machos = datos_iqr[datos_iqr["_GRUPO"] == "MACHOS"]
    hembras = datos_iqr[datos_iqr["_GRUPO"] == "HEMBRAS"]

    if len(machos) != 1:
        raise ValueError(
            f"Se esperaba una regresión IQR para Machos y se encontraron {len(machos)}")
    if len(hembras) != 1:
        raise ValueError(
            f"Se esperaba una regresión IQR para Hembras y se encontraron {len(hembras)}")

    machos = machos.iloc[0]
    hembras = hembras.iloc[0]
    return (float(machos["a"]), float(machos["b"]),
            float(hembras["a"]), float(hembras["b"]))

def leer_matriz_talla_edad(archivo, sexo, talla_min, talla_max, paso_talla,
                           edad_min, edad_max):
    libro = load_workbook(archivo, data_only=True)

    if sexo not in libro.sheetnames:
        raise ValueError(f"No existe la hoja '{sexo}' en {archivo.name}")

    hoja = libro[sexo]
    encabezados = {hoja.cell(1, col).value: col
                  for col in range(1, hoja.max_column + 1)}
    tallas = list(range(talla_min, talla_max + 1, paso_talla))
    edades = list(range(edad_min, edad_max + 1))
    datos = {talla: {edad: 0 for edad in edades} for talla in tallas}

    for fila in range(2, hoja.max_row + 1):
        grupo_talla = hoja.cell(fila, 1).value

        if grupo_talla in (None, "Total general"):
            continue

        try:
            grupo_talla = int(grupo_talla)
        except (TypeError, ValueError):
            continue

        if grupo_talla not in datos:
            continue

        for encabezado, columna in encabezados.items():
            if encabezado in ("GRUPO_TALLA", "Total general"):
                continue

            try:
                edad = int(encabezado)
            except (TypeError, ValueError):
                continue

            frecuencia = hoja.cell(fila, columna).value or 0
            edad_destino = min(edad_max, edad)

            if edad_destino in datos[grupo_talla]:
                datos[grupo_talla][edad_destino] += frecuencia

    return datos

def leer_frecuencias_colocap(archivo, hoja_nombre, paso_talla):
    libro = load_workbook(archivo, data_only=True)

    if hoja_nombre not in libro.sheetnames:
        raise ValueError(f"No se encontró la hoja {hoja_nombre}")

    hoja = libro[hoja_nombre]

    if any(hoja[celda].value is None for celda in ("L2", "M2", "N2")):
        raise ValueError("No se encontraron L2, M2 o N2 en el archivo Colocap")

    prop_machos = float(hoja["L2"].value)
    prop_hembras = float(hoja["M2"].value)
    total_ejemplares = float(hoja["N2"].value)
    n_machos = total_ejemplares * prop_machos
    n_hembras = total_ejemplares * prop_hembras
    frecuencias = {}

    for fila in range(3, hoja.max_row + 1):
        talla = hoja.cell(fila, 16).value
        prop_talla_m = hoja.cell(fila, 17).value or 0
        prop_talla_h = hoja.cell(fila, 18).value or 0

        try:
            talla = int(talla)
        except (TypeError, ValueError):
            continue

        marca_clase = talla + paso_talla // 2
        frecuencias[marca_clase] = {
            "Machos": n_machos * float(prop_talla_m),
            "Hembras": n_hembras * float(prop_talla_h)}

    return total_ejemplares, n_machos, n_hembras, frecuencias

# ============================================================
# BÚSQUEDA DE DESEMBARQUE
def normalizar_texto(valor):
    texto = str(valor or "").strip().upper()
    texto = "".join(caracter for caracter in unicodedata.normalize("NFD", texto)
                    if unicodedata.category(caracter) != "Mn")
    return " ".join(texto.split())

def buscar_desembarque(archivo, especie, flota):
    libro = load_workbook(archivo, data_only=True)
    especie_n = normalizar_texto(especie)
    flota_n = normalizar_texto(flota)
    encontrados = []

    for hoja in libro.worksheets:
        for fila in hoja.iter_rows():
            for celda in fila:
                if normalizar_texto(celda.value) != especie_n:
                    continue

                fila_especie = celda.row
                col_especie = celda.column
                fila_datos = None

                for fila_indice in range(fila_especie + 1,
                                         min(fila_especie + 16, hoja.max_row + 1)):
                    if normalizar_texto(hoja.cell(fila_indice, col_especie).value):
                        fila_datos = fila_indice
                        break

                if fila_datos is None:
                    continue

                fila_flota = None

                for fila_indice in range(fila_datos,
                                         min(fila_datos + 10, hoja.max_row + 1)):
                    if normalizar_texto(hoja.cell(fila_indice, col_especie).value) == flota_n:
                        fila_flota = fila_indice
                        break

                if fila_flota is None:
                    continue

                exactas = []
                parciales = []

                for col in range(col_especie + 1,
                                 min(col_especie + 10, hoja.max_column + 1)):
                    encabezados = [normalizar_texto(hoja.cell(fila_indice, col).value)
                                   for fila_indice in range(fila_especie, fila_datos)]

                    if "TOTAL" in encabezados:
                        exactas.append(col)
                    elif any("TOTAL" in encabezado for encabezado in encabezados):
                        parciales.append(col)

                for columnas in (exactas, parciales):
                    candidatos = []

                    for col in columnas:
                        valor = hoja.cell(fila_flota, col).value

                        if isinstance(valor, (int, float)) and not isinstance(valor, bool):
                            candidatos.append((col, float(valor)))

                    if candidatos:
                        _, valor = min(candidatos, key=lambda x: x[0])
                        encontrados.append(valor)
                        break

    if not encontrados:
        raise ValueError(
            f"No se encontró desembarque para '{especie}' / '{flota}'")

    unicos = []

    for valor in encontrados:
        if not any(abs(valor - existente) < 1e-9 for existente in unicos):
            unicos.append(valor)

    if len(unicos) > 1:
        valores = ", ".join(f"{valor:.6f}" for valor in unicos)
        raise ValueError(
            f"Se encontraron varios desembarques para '{especie}': {valores}")

    return unicos[0]
